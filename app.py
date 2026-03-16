import csv
import io
import json
from datetime import datetime, timezone

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from flask import (
    Flask,
    jsonify,
    render_template,
    request,
    send_file,
)
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from config import Config
from models import GroundModel, Stratum, db

# ── Column definitions for DGM tables ────────────────────────────────

DGM_HEADERS_ROW1 = [
    "Stratum", "No. of\nPoints (top)",
    "Top Level\n(m AOD)", "Top Level\n(m AOD)",
    "Bottom Level\n(m AOD)", "Bottom Level\n(m AOD)",
    "Top Level\n(m AOD)", "Bottom Level\n(m AOD)", "Thickness\n(m)",
]

DGM_HEADERS_ROW2 = [
    "", "", "Min.", "Max.", "Min.", "Max.", "Design", "Design", "Design",
]

STRATUM_KEYS = [
    "stratum_name", "num_points",
    "top_level_min", "top_level_max",
    "bottom_level_min", "bottom_level_max",
    "top_level_design", "bottom_level_design", "thickness_design",
]

CSV_HEADERS = [
    "Stratum", "No. of Points (top)",
    "Top Level Min (m AOD)", "Top Level Max (m AOD)",
    "Bottom Level Min (m AOD)", "Bottom Level Max (m AOD)",
    "Top Level Design (m AOD)", "Bottom Level Design (m AOD)",
    "Thickness Design (m)", "Ref", "Notes",
]


def _stratum_row(s):
    """Return display values for a stratum matching DGM table order."""
    return [
        s.stratum_name, s.num_points,
        s.top_level_min, s.top_level_max,
        s.bottom_level_min, s.bottom_level_max,
        s.top_level_design, s.bottom_level_design, s.thickness_design,
    ]


def create_app():
    app = Flask(__name__)
    app.config.from_object(Config)
    db.init_app(app)

    with app.app_context():
        db.create_all()

    # ── Page routes ──────────────────────────────────────────────────

    @app.route("/")
    def index():
        return render_template("index.html")

    @app.route("/model/<int:model_id>")
    def view_model(model_id):
        return render_template("model.html")

    @app.route("/shared/<share_id>")
    def view_shared(share_id):
        return render_template("model.html", shared=True)

    # ── API: list & create ───────────────────────────────────────────

    @app.route("/api/models", methods=["GET"])
    def list_models():
        models = GroundModel.query.order_by(GroundModel.updated_at.desc()).all()
        return jsonify([m.to_dict() for m in models])

    @app.route("/api/models", methods=["POST"])
    def create_model():
        data = request.get_json()
        if not data or not data.get("name"):
            return jsonify({"error": "Name is required"}), 400

        model = GroundModel(
            name=data["name"],
            title=data.get("title", ""),
            project=data.get("project", ""),
            location=data.get("location", ""),
            created_by=data.get("created_by", ""),
            rationale=data.get("rationale", ""),
            notes=data.get("notes", ""),
        )
        db.session.add(model)
        db.session.commit()
        return jsonify(model.to_dict()), 201

    # ── API: single model CRUD ───────────────────────────────────────

    @app.route("/api/models/<int:model_id>", methods=["GET"])
    def get_model(model_id):
        model = db.session.get(GroundModel, model_id)
        if not model:
            return jsonify({"error": "Not found"}), 404
        return jsonify(model.to_dict())

    @app.route("/api/models/<int:model_id>", methods=["PUT"])
    def update_model(model_id):
        model = db.session.get(GroundModel, model_id)
        if not model:
            return jsonify({"error": "Not found"}), 404
        data = request.get_json()
        for field in ("name", "title", "project", "location", "created_by", "rationale", "notes"):
            if field in data:
                setattr(model, field, data[field])
        model.updated_at = datetime.now(timezone.utc)
        db.session.commit()
        return jsonify(model.to_dict())

    @app.route("/api/models/<int:model_id>", methods=["DELETE"])
    def delete_model(model_id):
        model = db.session.get(GroundModel, model_id)
        if not model:
            return jsonify({"error": "Not found"}), 404
        db.session.delete(model)
        db.session.commit()
        return jsonify({"ok": True})

    # ── API: shared access ───────────────────────────────────────────

    @app.route("/api/shared/<share_id>", methods=["GET"])
    def get_shared_model(share_id):
        model = GroundModel.query.filter_by(share_id=share_id).first()
        if not model:
            return jsonify({"error": "Not found"}), 404
        return jsonify(model.to_dict())

    # ── API: strata ──────────────────────────────────────────────────

    @app.route("/api/models/<int:model_id>/strata", methods=["POST"])
    def add_stratum(model_id):
        model = db.session.get(GroundModel, model_id)
        if not model:
            return jsonify({"error": "Model not found"}), 404
        data = request.get_json()
        if not data or not data.get("stratum_name"):
            return jsonify({"error": "stratum_name is required"}), 400

        max_order = max((s.sort_order for s in model.strata), default=-1)
        stratum = Stratum(
            ground_model_id=model_id,
            sort_order=max_order + 1,
            stratum_name=data["stratum_name"],
            ref_number=data.get("ref_number", ""),
            num_points=data.get("num_points", ""),
            top_level_min=data.get("top_level_min", ""),
            top_level_max=data.get("top_level_max", ""),
            bottom_level_min=data.get("bottom_level_min", ""),
            bottom_level_max=data.get("bottom_level_max", ""),
            top_level_design=data.get("top_level_design", ""),
            bottom_level_design=data.get("bottom_level_design", ""),
            thickness_design=data.get("thickness_design", ""),
            stratum_notes=data.get("stratum_notes", ""),
        )
        db.session.add(stratum)
        model.updated_at = datetime.now(timezone.utc)
        db.session.commit()
        return jsonify(stratum.to_dict()), 201

    @app.route("/api/strata/<int:stratum_id>", methods=["PUT"])
    def update_stratum(stratum_id):
        stratum = db.session.get(Stratum, stratum_id)
        if not stratum:
            return jsonify({"error": "Not found"}), 404
        data = request.get_json()
        for field in (
            "stratum_name", "ref_number", "num_points",
            "top_level_min", "top_level_max",
            "bottom_level_min", "bottom_level_max",
            "top_level_design", "bottom_level_design", "thickness_design",
            "stratum_notes", "sort_order",
        ):
            if field in data:
                val = data[field]
                if field == "sort_order":
                    val = int(val)
                setattr(stratum, field, val)
        stratum.ground_model.updated_at = datetime.now(timezone.utc)
        db.session.commit()
        return jsonify(stratum.to_dict())

    @app.route("/api/strata/<int:stratum_id>", methods=["DELETE"])
    def delete_stratum(stratum_id):
        stratum = db.session.get(Stratum, stratum_id)
        if not stratum:
            return jsonify({"error": "Not found"}), 404
        stratum.ground_model.updated_at = datetime.now(timezone.utc)
        db.session.delete(stratum)
        db.session.commit()
        return jsonify({"ok": True})

    # ── Export: JSON ─────────────────────────────────────────────────

    @app.route("/api/models/<int:model_id>/export/json")
    def export_json(model_id):
        model = db.session.get(GroundModel, model_id)
        if not model:
            return jsonify({"error": "Not found"}), 404
        data = model.to_dict()
        data.pop("id")
        for s in data["strata"]:
            s.pop("id")
        buf = io.BytesIO(json.dumps(data, indent=2).encode("utf-8"))
        filename = f"{model.name.replace(' ', '_')}_ground_model.json"
        return send_file(buf, mimetype="application/json", as_attachment=True, download_name=filename)

    # ── Export: CSV ──────────────────────────────────────────────────

    @app.route("/api/models/<int:model_id>/export/csv")
    def export_csv(model_id):
        model = db.session.get(GroundModel, model_id)
        if not model:
            return jsonify({"error": "Not found"}), 404
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(CSV_HEADERS)
        for s in model.strata:
            writer.writerow([
                s.stratum_name, s.num_points,
                s.top_level_min, s.top_level_max,
                s.bottom_level_min, s.bottom_level_max,
                s.top_level_design, s.bottom_level_design, s.thickness_design,
                s.ref_number, s.stratum_notes,
            ])
        output = io.BytesIO(buf.getvalue().encode("utf-8"))
        filename = f"{model.name.replace(' ', '_')}_ground_model.csv"
        return send_file(output, mimetype="text/csv", as_attachment=True, download_name=filename)

    # ── Export: PDF ──────────────────────────────────────────────────

    @app.route("/api/models/<int:model_id>/export/pdf")
    def export_pdf(model_id):
        model = db.session.get(GroundModel, model_id)
        if not model:
            return jsonify({"error": "Not found"}), 404

        buf = io.BytesIO()
        doc = SimpleDocTemplate(
            buf, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm
        )
        styles = getSampleStyleSheet()
        elements = []

        # Title
        table_label = f"{model.name}"
        if model.title:
            table_label += f" \u2013 {model.title}"
        elements.append(Paragraph(table_label, styles["Title"]))

        # Info
        for label, val in [("Project", model.project), ("Location", model.location),
                           ("Created by", model.created_by)]:
            if val:
                elements.append(Paragraph(f"<b>{label}:</b> {val}", styles["Normal"]))
        elements.append(Spacer(1, 6 * mm))

        # DGM table with merged header rows
        table_data = [DGM_HEADERS_ROW1, DGM_HEADERS_ROW2]
        for s in model.strata:
            table_data.append(_stratum_row(s))

        col_widths = [120, 55, 55, 55, 55, 55, 75, 75, 65]
        table = Table(table_data, colWidths=col_widths, repeatRows=2)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 1), colors.HexColor("#2c3e50")),
            ("TEXTCOLOR", (0, 0), (-1, 1), colors.white),
            ("FONTNAME", (0, 0), (-1, 1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 1), 7),
            ("FONTSIZE", (0, 2), (-1, -1), 7),
            ("ALIGN", (1, 0), (-1, -1), "CENTER"),
            ("ALIGN", (0, 0), (0, -1), "LEFT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("ROWBACKGROUNDS", (0, 2), (-1, -1), [colors.white, colors.HexColor("#ecf0f1")]),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            # Merge header cells for Stratum and No. of Points across rows 0-1
            ("SPAN", (0, 0), (0, 1)),
            ("SPAN", (1, 0), (1, 1)),
            # Merge Top Level across cols 2-3 in row 0
            ("SPAN", (2, 0), (3, 0)),
            # Merge Bottom Level across cols 4-5 in row 0
            ("SPAN", (4, 0), (5, 0)),
        ]))
        elements.append(table)

        # Notes table
        notes_strata = [s for s in model.strata if s.ref_number and s.stratum_notes]
        if notes_strata:
            elements.append(Spacer(1, 4 * mm))
            notes_data = []
            for s in notes_strata:
                notes_data.append([s.ref_number, s.stratum_name.split("[")[0].strip(), s.stratum_notes])
            notes_table = Table(notes_data, colWidths=[30, 100, 480])
            notes_table.setStyle(TableStyle([
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]))
            elements.append(notes_table)

        # Rationale
        if model.rationale:
            elements.append(Spacer(1, 6 * mm))
            elements.append(Paragraph("<b>Design Ground Model Rationale</b>", styles["Normal"]))
            elements.append(Spacer(1, 2 * mm))
            elements.append(Paragraph(model.rationale, styles["Normal"]))

        doc.build(elements)
        buf.seek(0)
        filename = f"{model.name.replace(' ', '_')}_ground_model.pdf"
        return send_file(buf, mimetype="application/pdf", as_attachment=True, download_name=filename)

    # ── Export: Word (DOCX) ──────────────────────────────────────────

    @app.route("/api/models/<int:model_id>/export/docx")
    def export_docx(model_id):
        model = db.session.get(GroundModel, model_id)
        if not model:
            return jsonify({"error": "Not found"}), 404

        doc = Document()

        # Set landscape
        section = doc.sections[0]
        new_width, new_height = section.page_height, section.page_width
        section.page_width = new_width
        section.page_height = new_height
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

        # Table title (matching template style: "Table 3-X – DGM_ref – Title")
        table_label = model.name
        if model.title:
            table_label += f" \u2013 {model.title}"
        title_para = doc.add_paragraph()
        title_para.style = doc.styles["Normal"]
        run = title_para.add_run(table_label)
        run.bold = True
        run.font.size = Pt(10)

        # Project info
        if model.project:
            p = doc.add_paragraph()
            p.add_run("Project: ").bold = True
            p.add_run(model.project)
        if model.location:
            p = doc.add_paragraph()
            p.add_run("Location: ").bold = True
            p.add_run(model.location)
        if model.created_by:
            p = doc.add_paragraph()
            p.add_run("Created by: ").bold = True
            p.add_run(model.created_by)

        # ── Main DGM table ───────────────────────────────────────────
        num_strata = len(model.strata)
        table = doc.add_table(rows=2 + num_strata, cols=9)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True

        # Header row 1
        h1_texts = [
            "Stratum", "No. of\nPoints (top)",
            "Top Level\n(m AOD)", "",
            "Bottom Level\n(m AOD)", "",
            "Top Level\n(m AOD)", "Bottom Level\n(m AOD)", "Thickness\n(m)",
        ]
        for i, text in enumerate(h1_texts):
            cell = table.cell(0, i)
            _set_cell_text(cell, text, bold=True, size=8, align="center")
            _shade_cell(cell, "2C3E50")
            _set_cell_text_color(cell, "FFFFFF")

        # Header row 2
        h2_texts = ["", "", "Min.", "Max.", "Min.", "Max.", "Design", "Design", "Design"]
        for i, text in enumerate(h2_texts):
            cell = table.cell(1, i)
            _set_cell_text(cell, text, bold=True, size=8, align="center")
            _shade_cell(cell, "2C3E50")
            _set_cell_text_color(cell, "FFFFFF")

        # Merge header cells: Stratum spans rows 0-1
        table.cell(0, 0).merge(table.cell(1, 0))
        # No. of Points spans rows 0-1
        table.cell(0, 1).merge(table.cell(1, 1))
        # Top Level spans cols 2-3 in row 0
        table.cell(0, 2).merge(table.cell(0, 3))
        # Bottom Level spans cols 4-5 in row 0
        table.cell(0, 4).merge(table.cell(0, 5))

        # Data rows
        for row_idx, s in enumerate(model.strata):
            row_num = row_idx + 2
            vals = _stratum_row(s)
            for col_idx, val in enumerate(vals):
                cell = table.cell(row_num, col_idx)
                align = "left" if col_idx == 0 else "center"
                _set_cell_text(cell, str(val), size=8, align=align)
                # Alternate row shading
                if row_idx % 2 == 1:
                    _shade_cell(cell, "ECF0F1")

        # ── Notes table ──────────────────────────────────────────────
        notes_strata = [s for s in model.strata if s.ref_number and s.stratum_notes]
        if notes_strata:
            doc.add_paragraph()  # spacer
            notes_table = doc.add_table(rows=len(notes_strata), cols=3)
            notes_table.style = "Table Grid"
            notes_table.alignment = WD_TABLE_ALIGNMENT.CENTER

            for i, s in enumerate(notes_strata):
                _set_cell_text(notes_table.cell(i, 0), s.ref_number, size=8)
                # Strip ref from name for notes column
                clean_name = s.stratum_name.split("[")[0].strip()
                _set_cell_text(notes_table.cell(i, 1), clean_name, bold=True, size=8)
                _set_cell_text(notes_table.cell(i, 2), s.stratum_notes, size=8)

        # ── Rationale ────────────────────────────────────────────────
        if model.rationale:
            doc.add_paragraph()
            p = doc.add_paragraph()
            run = p.add_run("Design Ground Model Rationale")
            run.bold = True
            run.font.size = Pt(9)
            p = doc.add_paragraph(model.rationale)
            for run in p.runs:
                run.font.size = Pt(9)

        # Model notes
        if model.notes:
            doc.add_paragraph()
            p = doc.add_paragraph()
            run = p.add_run("Notes: ")
            run.bold = True
            run.font.size = Pt(9)
            p.add_run(model.notes).font.size = Pt(9)

        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        filename = f"{model.name.replace(' ', '_')}_ground_model.docx"
        return send_file(
            buf,
            mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            as_attachment=True,
            download_name=filename,
        )

    # ── Import: JSON / CSV ───────────────────────────────────────────

    @app.route("/api/models/import", methods=["POST"])
    def import_model():
        if "file" not in request.files:
            return jsonify({"error": "No file provided"}), 400
        file = request.files["file"]
        filename = file.filename or ""

        try:
            if filename.endswith(".json"):
                data = json.load(file)
            elif filename.endswith(".csv"):
                reader = csv.DictReader(io.TextIOWrapper(file, encoding="utf-8"))
                data = {
                    "name": filename.rsplit(".", 1)[0].replace("_", " "),
                    "strata": [],
                }
                for i, row in enumerate(reader):
                    data["strata"].append({
                        "sort_order": i,
                        "stratum_name": row.get("Stratum", ""),
                        "num_points": row.get("No. of Points (top)", ""),
                        "top_level_min": row.get("Top Level Min (m AOD)", ""),
                        "top_level_max": row.get("Top Level Max (m AOD)", ""),
                        "bottom_level_min": row.get("Bottom Level Min (m AOD)", ""),
                        "bottom_level_max": row.get("Bottom Level Max (m AOD)", ""),
                        "top_level_design": row.get("Top Level Design (m AOD)", ""),
                        "bottom_level_design": row.get("Bottom Level Design (m AOD)", ""),
                        "thickness_design": row.get("Thickness Design (m)", ""),
                        "ref_number": row.get("Ref", ""),
                        "stratum_notes": row.get("Notes", ""),
                    })
            else:
                return jsonify({"error": "Unsupported file type. Use .json or .csv"}), 400
        except Exception as e:
            return jsonify({"error": f"Failed to parse file: {e}"}), 400

        model = GroundModel(
            name=data.get("name", "Imported Model"),
            title=data.get("title", ""),
            project=data.get("project", ""),
            location=data.get("location", ""),
            created_by=data.get("created_by", ""),
            rationale=data.get("rationale", ""),
            notes=data.get("notes", ""),
        )
        db.session.add(model)
        db.session.flush()

        for i, sd in enumerate(data.get("strata", [])):
            stratum = Stratum(
                ground_model_id=model.id,
                sort_order=sd.get("sort_order", i),
                stratum_name=sd.get("stratum_name", ""),
                ref_number=sd.get("ref_number", ""),
                num_points=sd.get("num_points", ""),
                top_level_min=sd.get("top_level_min", ""),
                top_level_max=sd.get("top_level_max", ""),
                bottom_level_min=sd.get("bottom_level_min", ""),
                bottom_level_max=sd.get("bottom_level_max", ""),
                top_level_design=sd.get("top_level_design", ""),
                bottom_level_design=sd.get("bottom_level_design", ""),
                thickness_design=sd.get("thickness_design", ""),
                stratum_notes=sd.get("stratum_notes", ""),
            )
            db.session.add(stratum)

        db.session.commit()
        return jsonify(model.to_dict()), 201

    return app


# ── DOCX helper functions ────────────────────────────────────────────

def _set_cell_text(cell, text, bold=False, size=9, align="left"):
    """Set cell text with formatting, clearing existing content."""
    cell.text = ""
    para = cell.paragraphs[0]
    run = para.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    if align == "center":
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        para.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _shade_cell(cell, color_hex):
    """Apply background shading to a cell."""
    tc_pr = cell._element.get_or_add_tcPr()
    shading = tc_pr.makeelement(
        qn("w:shd"),
        {qn("w:fill"): color_hex, qn("w:val"): "clear"},
    )
    tc_pr.append(shading)


def _set_cell_text_color(cell, color_hex):
    """Set text color for all runs in a cell."""
    r, g, b = int(color_hex[:2], 16), int(color_hex[2:4], 16), int(color_hex[4:], 16)
    for para in cell.paragraphs:
        for run in para.runs:
            run.font.color.rgb = RGBColor(r, g, b)


if __name__ == "__main__":
    app = create_app()
    app.run(debug=True, host="0.0.0.0", port=5000)
